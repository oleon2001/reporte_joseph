"""
Generación de la hoja de cálculo: replica EXACTAMENTE el diseño y el análisis
de "Detalles de Tr's atrapadas DD-MM-YYYY.xlsx" del equipo de continuidad.

Diseño observado en el archivo de referencia:
  - Sin relleno/bold. Encabezados de sección en columna D combinados D:G.
  - Guatemala:
      A1='StellarDocID'  B1='Observaciones'  B2='Sin observaciones.'
      Secciones en D (merged D:G), cada 4 filas:
        'Sin cliente en VENTAS.MAECEC o Sin documento en VENTAS.OPEVTSC'
        'Sin cliente en CR400'
        'Transacciones con estatus en blanco en sus detalles'
        'Devoluciones sin detalles'
        "TR's Pendientes por imprimir"  (Tienda/POS/TR/Mensaje)
  - Salvador: header A1-D1 centrado; Tienda/POS/TR EPA/Observaciones.
  - Venezuela: header A1-B1; StellarDocID/Observaciones.
  - La columna Observaciones tiene LISTA DESPLEGABLE (data validation) con las
    opciones fijas definidas en config.OBSERVACIONES.
"""
import logging
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config

log = logging.getLogger(__name__)

_WIDTHS_GT = {"A": 19.13, "B": 77.88, "C": 15.75, "D": 18.75, "G": 14.13}
_WIDTHS_SV = {"A": 9.25, "B": 7.5, "C": 12.25, "D": 66.0}
_WIDTHS_VE = {"A": 19.13, "B": 106.0, "C": 15.75, "D": 38.13}

_CENTER = Alignment(horizontal="center")


def _num(value):
    """El archivo real guarda Tienda/POS/TR como número."""
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def _apply_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _add_dropdown(ws, options, col_letter, start_row, end_row, list_col=None):
    """Agrega la lista desplegable de Observaciones.

    La lista inline tiene un límite de 255 caracteres, así que las opciones se
    guardan en una hoja oculta "Listas" y la validación apunta a ese rango.
    """
    options = [o.rstrip(",") for o in options if o]
    list_sheet = ws.parent["Listas"] if "Listas" in ws.parent.sheetnames else None
    if list_sheet is None:
        list_sheet = ws.parent.create_sheet("Listas")
        list_sheet.sheet_state = "hidden"
    col_idx = list_sheet.max_column + 1
    col_letter_ref = get_column_letter(col_idx)
    for i, opt in enumerate(options, start=1):
        list_sheet.cell(row=i, column=col_idx, value=opt)
    list_sheet.column_dimensions[col_letter_ref].hidden = True
    ref = f"Listas!${col_letter_ref}$1:${col_letter_ref}${len(options)}"
    dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    last = max(end_row, start_row)
    dv.add(f"{col_letter}{start_row}:{col_letter}{last}")


def _section(ws, start_row, titulo, headers, data_rows, rows_of_data=2):
    """Escribe una sección en col D (merged D:G), separando cada 4 filas."""
    ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=7)
    ws.cell(row=start_row, column=4).value = titulo
    hdr_row = start_row + 1
    for i, h in enumerate(headers):
        ws.cell(row=hdr_row, column=4 + i).value = h
    for j in range(rows_of_data):
        r = hdr_row + 1 + j
        if j < len(data_rows):
            row_data = data_rows[j]
            for i in range(len(headers)):
                ws.cell(row=r, column=4 + i).value = row_data[i] if i < len(row_data) else None
    return start_row + 4


# ---------------------------------------------------------------------------
# Por país
# ---------------------------------------------------------------------------
def _sheet_guatemala(ws, rows):
    def take(origen):
        # Coincidencia por prefijo: los orígenes por tienda llevan sufijo
        # (ej. "MA_PAGOS pendiente tienda 5").
        return [r for r in rows if r["Origen"].startswith(origen)]

    # Bloque Cabtra (A/B) con desplegable en B
    ws["A1"] = "StellarDocID"
    ws["B1"] = "Observaciones"
    cabtra = take("Cabtra GT")
    if cabtra:
        for i, r in enumerate(cabtra):
            ws.cell(row=2 + i, column=1).value = r["StellarDocID"]
            ws.cell(row=2 + i, column=2).value = r["Observaciones"]
        _add_dropdown(ws, config.OBSERVACIONES["GT"], "B", 2, 1 + len(cabtra))
    else:
        ws["B2"] = "Sin observaciones."
        _add_dropdown(ws, config.OBSERVACIONES["GT"], "B", 2, 2)

    row = _section(
        ws, 1,
        "Sin cliente en VENTAS.MAECEC o Sin documento en VENTAS.OPEVTSC",
        ["Tienda", "POS", "TR EPA"],
        [[_num(r["Tienda"]), _num(r["POS/Caja"]), _num(r["TR EPA"])]
         for r in take("AS400: no llegó cliente a MAECEC") + take("AS400: no llegó a OPVTSC")],
    )
    row = _section(
        ws, row,
        "Sin cliente en CR400",
        ["Tienda", "POS", "TR EPA", "CODCTE"],
        [[_num(r["Tienda"]), _num(r["POS/Caja"]), _num(r["TR EPA"]), r.get("CODCTE", "")]
         for r in take("AS400: no llegó cliente a CR400")],
    )
    row = _section(
        ws, row,
        "Transacciones con estatus en blanco en sus detalles",
        ["Tienda", "POS", "TR EPA"],
        [[_num(r["Tienda"]), _num(r["POS/Caja"]), _num(r["TR EPA"])]
         for r in take("AS400: estatus en blanco")],
    )
    row = _section(
        ws, row,
        "Devoluciones sin detalles",
        ["Tienda", "POS", "TR EPA"],
        [[_num(r["Tienda"]), _num(r["POS/Caja"]), _num(r["TR EPA"])]
         for r in take("AS400: devolución sin detalle")],
    )
    _section(
        ws, row,
        "TR's Pendientes por imprimir",
        ["Tienda", "POS", "TR", "Mensaje"],
        [[_num(r["Tienda"]), _num(r["POS/Caja"]), _num(r["TR EPA"]), r["Observaciones"]]
         for r in take("MA_PAGOS pendiente")],
    )
    _apply_widths(ws, _WIDTHS_GT)


def _sheet_salvador(ws, rows):
    for i, h in enumerate(["Tienda", "POS", "TR EPA", "Observaciones"]):
        cell = ws.cell(row=1, column=1 + i)
        cell.value = h
        cell.alignment = _CENTER
    for r, item in enumerate(rows, start=2):
        ws.cell(row=r, column=1).value = _num(item["Tienda"])
        ws.cell(row=r, column=2).value = _num(item["POS/Caja"])
        ws.cell(row=r, column=3).value = _num(item["TR EPA"])
        ws.cell(row=r, column=4).value = item["Observaciones"]
    last = max(2, 1 + len(rows))
    _add_dropdown(ws, config.OBSERVACIONES["SV"], "D", 2, last)
    _apply_widths(ws, _WIDTHS_SV)


def _sheet_venezuela(ws, rows):
    ws["A1"] = "StellarDocID"
    ws["B1"] = "Observaciones"
    for r, item in enumerate(rows, start=2):
        ws.cell(row=r, column=1).value = item["StellarDocID"]
        ws.cell(row=r, column=2).value = item["Observaciones"]
    last = max(2, 1 + len(rows))
    _add_dropdown(ws, config.OBSERVACIONES["VE"], "B", 2, last)
    _apply_widths(ws, _WIDTHS_VE)


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------
def build_workbook(rows_by_country, out_dir, date_str):
    os.makedirs(out_dir, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    builders = {"GT": _sheet_guatemala, "SV": _sheet_salvador, "VE": _sheet_venezuela}
    for country in config.SHEET_ORDER:
        meta = config.COUNTRIES[country]
        ws = wb.create_sheet(meta["sheet"])
        builders[country](ws, rows_by_country.get(country, []))
    path = os.path.join(out_dir, f"Detalles de Tr's atrapadas {date_str}.xlsx")
    wb.active = 0  # la primera pestaña de datos queda activa
    wb.save(path)
    log.info("Hoja de cálculo generada: %s", path)
    return path


if __name__ == "__main__":
    import tempfile

    sample = {"GT": [{"Tienda": "3", "POS/Caja": "52", "TR EPA": "52023419",
                      "StellarDocID": "X", "Origen": "MA_PAGOS pendiente",
                      "Observaciones": "FAIL - 82-ERROR"}],
              "SV": [{"Tienda": "3", "POS/Caja": "20", "TR EPA": "666745",
                      "Observaciones": "Se desconoce el motivo ya que la caja no está encendida."}],
              "VE": [{"StellarDocID": "16DEV41012504",
                      "Observaciones": "Falta en VENTASSTORE pero si se encuentra en VAD20."}]}
    with tempfile.TemporaryDirectory() as tmp:
        out = build_workbook(sample, tmp, "01-01-2026")
        from openpyxl import load_workbook

        wb = load_workbook(out)
        data_sheets = [s for s in wb.sheetnames if s != "Listas"]
        assert data_sheets == ["Guatemala", "Salvador", "Venezuela"], wb.sheetnames
        assert "Listas" in wb.sheetnames and wb["Listas"].sheet_state == "hidden"
        g = wb["Guatemala"]
        assert g["A1"].value == "StellarDocID" and g["B1"].value == "Observaciones"
        assert g["B2"].value == "Sin observaciones."
        assert g["D1"].value.startswith("Sin cliente en VENTAS.MAECEC")
        assert g["D17"].value == "TR's Pendientes por imprimir"
        assert len(g.data_validations.dataValidation) >= 1
        v = wb["Venezuela"]
        assert v["A1"].value == "StellarDocID"
        assert len(v.data_validations.dataValidation) >= 1
        s = wb["Salvador"]
        assert len(s.data_validations.dataValidation) >= 1
    print("excel_builder OK")
